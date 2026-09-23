import json
import zipfile

import pytest
from openpyxl import load_workbook

from geopol import excel_export
from geopol.excel_export import HEADER_ROW, write_xlsx


COLUMNS = (
    "id",
    "complaint_id",
    "location_original",
    "location_normalized",
    "ubigeo",
    "resolution",
    "method",
    "precision",
    "evidence_band",
    "product",
    "latitude",
    "longitude",
    "reason",
    "revision",
    "source_row_count",
    "review_status",
    "review_bucket",
    "geometry",
)


def metadata(**overrides):
    return {
        "profile": "locations",
        "source_columns": ["denuncia", "direccion", "__extra_columns__"],
        "run_name": "Muestra sintética",
        "filename": "muestra.xlsx",
        "expected_rows": 1,
        "snapshot_at": "2026-09-23T14:00:00Z",
        "rules_version": "2026.3",
        "reference": {"name": "Catálogo sintético", "version": "2026.1"},
        **overrides,
    }


def row(**snapshot):
    return {
        "ordinal": 1,
        "issue": None,
        "raw": {"denuncia": "000123", "direccion": "Calle Prueba 12"},
        "snapshot": {
            "id": "00000000-0000-4000-8000-000000000001",
            "complaint_id": "000123",
            "location_original": "Calle Prueba 12",
            "location_normalized": "CALLE PRUEBA 12",
            "ubigeo": "010101",
            "resolution": "ACEPTADO_AUTOMATICO",
            "method": "PUERTA_EXACTA",
            "precision": "PUERTA",
            "evidence_band": "ALTA",
            "product": "PUNTO",
            "latitude": -12.123456789,
            "longitude": -77.01234567,
            "reason": "Coincidencia corroborada",
            "revision": 1,
            "source_row_count": 2,
            "review_status": "CLOSED",
            "review_bucket": "none",
            "geometry": {"type": "Point", "coordinates": [-77.01234567, -12.123456789]},
            **snapshot,
        },
    }


def result_cells(book, stats, number=HEADER_ROW + 1):
    return {
        column["key"]: book["Resultados"].cell(number, index)
        for index, column in enumerate(stats["columns"], 1)
    }


def test_readable_workbook_preserves_snapshot_types_and_geographic_precision(tmp_path):
    path = tmp_path / "result.xlsx"
    original = row()
    stats = write_xlsx(path, iter([original]), columns=COLUMNS, metadata=metadata())
    book = load_workbook(path)
    sheet = book["Resultados"]
    cells = result_cells(book, stats)
    assert book.sheetnames == ["Resultados"]
    assert set(cells) == set(COLUMNS)
    assert stats["row_count"] == 1
    assert sheet.freeze_panes == "C8"
    assert sheet.auto_filter.ref == "A7:R8"
    assert not sheet.sheet_view.showGridLines
    assert sheet["A7"].fill.fgColor.rgb == "0012426B"
    assert sheet["A7"].font.color.rgb == "00FFFFFF"
    assert sheet["A1"].font.name == "Arial"
    assert sheet["A1"].font.sz == 14
    assert cells["resolution"].fill.fgColor.rgb == "00E5F2EC"
    assert sheet.sheet_format.defaultRowHeight == 42
    assert cells["complaint_id"].value == "000123"
    assert cells["ubigeo"].value == "010101"
    assert cells["ubigeo"].data_type == "s"
    assert cells["latitude"].data_type == "n"
    assert cells["latitude"].value == original["snapshot"]["latitude"]
    assert cells["longitude"].number_format == "0.00000000"
    assert json.loads(cells["geometry"].value) == original["snapshot"]["geometry"]
    assert sheet.column_dimensions[cells["id"].column_letter].hidden
    assert sheet.column_dimensions[cells["id"].column_letter].outlineLevel == 1
    assert "Catálogo sintético" in sheet["A4"].value
    assert stats["overflow"] == {"cells": 0, "segments": 0}
    book.close()


def test_source_rows_keep_duplicates_issues_original_order_and_no_inferred_dates(tmp_path):
    rows = [row(), row(latitude=None, longitude=None, product="AREA", precision="VIA")]
    rows[1].update(
        ordinal=2,
        issue="Incidencia sintética",
        raw={"=encabezado": "=1+2", "fecha": "01/02/2026", "ubigeo": 10101, "numero": 12345678901234567890},
    )
    columns = ["=encabezado", "fecha", "ubigeo", "numero", "__extra_columns__"]
    path = tmp_path / "originals.xlsx"
    stats = write_xlsx(
        path,
        rows,
        columns=COLUMNS,
        metadata=metadata(profile="source_rows", source_columns=columns, expected_rows=2),
    )
    book = load_workbook(path)
    assert book.sheetnames == ["Resultados", "Datos originales"]
    sheet = book["Datos originales"]
    assert [cell.value for cell in sheet[7]] == [
        "Ordinal de origen",
        *columns,
        "ID de ubicación",
        "Incidencia de origen",
    ]
    assert sheet["B7"].data_type == "s"
    assert sheet["B9"].value == "=1+2"
    assert sheet["B9"].data_type == "s"
    assert sheet["C9"].value == "01/02/2026"
    assert sheet["C9"].data_type == "s"
    assert sheet["D9"].value == 10101
    assert sheet["D9"].data_type == "n"
    assert sheet["E9"].value == "12345678901234567890"
    assert sheet["H9"].value == "Incidencia sintética"
    assert result_cells(book, stats, 9)["latitude"].value is None
    assert result_cells(book, stats, 9)["longitude"].value is None
    assert result_cells(book, stats)["id"].value == result_cells(book, stats, 9)["id"].value
    assert stats["sheets"][0]["data_rows"] == stats["sheets"][1]["data_rows"] == 2
    book.close()


@pytest.mark.parametrize(
    "value",
    [
        '=HYPERLINK("https://invalid.test","x")',
        "+123",
        "-12",
        "@SUM(A1:A2)",
        "\t =1+2",
        "00123456789012345678",
    ],
)
def test_untrusted_strings_are_literals_without_mutation(tmp_path, value):
    path = tmp_path / "literal.xlsx"
    stats = write_xlsx(path, [row(complaint_id=value, reason=value)], columns=COLUMNS, metadata=metadata())
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in xml
    book = load_workbook(path, data_only=False)
    cells = result_cells(book, stats)
    assert cells["complaint_id"].value == value
    assert cells["reason"].value == value
    assert cells["reason"].data_type == "s"
    book.close()


def test_long_geometry_and_source_text_round_trip_without_truncation(tmp_path):
    path = tmp_path / "long.xlsx"
    geometry = {"type": "LineString", "coordinates": [[-77.123456, -12.123456] for _ in range(4000)]}
    long_text = "=Texto original á\n" + "\U0001f600" * 20000
    original = row(geometry=geometry, reason="E" * 33000)
    original["raw"]["direccion"] = long_text
    stats = write_xlsx(path, [original], columns=COLUMNS, metadata=metadata(profile="source_rows"))
    assert stats["overflow"]["cells"] == 3
    assert stats["overflow"]["segments"] > 3
    book = load_workbook(path)
    assert book.sheetnames == ["Resultados", "Datos originales", "Textos extensos"]
    recovered = {}
    keys = {}
    for token, source_sheet, source_row, key, part, count, text in book["Textos extensos"].iter_rows(
        min_row=4, values_only=True
    ):
        assert len(text.encode("utf-16-le")) // 2 <= 30000
        recovered.setdefault(token, []).append((part, text))
        keys[(source_sheet, source_row, key)] = token
        assert 1 <= part <= count
    values = {key: "".join(value for _, value in sorted(recovered[token])) for key, token in keys.items()}
    assert json.loads(values[("Resultados", 8, "geometry")]) == geometry
    assert values[("Resultados", 8, "reason")] == "E" * 33000
    assert values[("Datos originales", 8, "origen:direccion")] == long_text
    assert "Texto completo" in result_cells(book, stats)["geometry"].value
    book.close()


def test_cell_at_excel_limit_needs_no_overflow_sheet(tmp_path):
    path = tmp_path / "limit.xlsx"
    stats = write_xlsx(path, [row(reason="x" * 32767)], columns=COLUMNS, metadata=metadata())
    book = load_workbook(path)
    assert len(result_cells(book, stats)["reason"].value) == 32767
    assert "Textos extensos" not in book.sheetnames
    book.close()


def test_empty_workbook_has_headers_and_filters_and_no_phantom_records(tmp_path):
    path = tmp_path / "empty.xlsx"
    stats = write_xlsx(path, iter([]), columns=COLUMNS, metadata=metadata(expected_rows=0))
    book = load_workbook(path)
    assert stats["row_count"] == 0
    assert book["Resultados"].max_row == HEADER_ROW
    assert book["Resultados"].auto_filter.ref == "A7:R7"
    book.close()


@pytest.mark.parametrize("value", ["control\x00", "control\x0b", "control\ud800", "control\ufffe"])
def test_invalid_xml_fails_with_csv_guidance_instead_of_dropping_text(tmp_path, value):
    path = tmp_path / "bad.xlsx"
    with pytest.raises(ValueError, match="CSV"):
        write_xlsx(path, [row(reason=value)], columns=COLUMNS, metadata=metadata())
    assert not path.exists()


def test_cardinality_and_excel_limits_fail_clearly(tmp_path, monkeypatch):
    path = tmp_path / "bad.xlsx"
    with pytest.raises(ValueError, match="cardinalidad"):
        write_xlsx(path, iter([]), columns=COLUMNS, metadata=metadata())
    monkeypatch.setattr(excel_export, "MAX_ROWS", HEADER_ROW + 1)
    with pytest.raises(ValueError, match="CSV"):
        write_xlsx(path, [row(), row()], columns=COLUMNS, metadata=metadata(expected_rows=None))
    with pytest.raises(ValueError, match="CSV"):
        write_xlsx(path, [], columns=COLUMNS, metadata=metadata(expected_rows=2))
    monkeypatch.setattr(excel_export, "MAX_COLUMNS", 19)
    with pytest.raises(ValueError, match="CSV"):
        write_xlsx(
            path,
            [],
            columns=COLUMNS,
            metadata=metadata(profile="source_rows", source_columns=[str(i) for i in range(17)]),
        )


def test_writer_consumes_one_pass_without_collecting_records(tmp_path):
    class OnePass:
        def __init__(self):
            self.started = False

        def __iter__(self):
            assert not self.started
            self.started = True
            for index in range(500):
                yield row(complaint_id=str(index))

    path = tmp_path / "stream.xlsx"
    stats = write_xlsx(path, OnePass(), columns=COLUMNS, metadata=metadata(expected_rows=500))
    book = load_workbook(path, read_only=True)
    assert sum(1 for _ in book["Resultados"].iter_rows(min_row=8)) == 500
    assert stats["row_count"] == 500
    book.close()


def test_source_header_collisions_preserve_names_with_distinct_trace_columns(tmp_path):
    path = tmp_path / "collisions.xlsx"
    names = ["Ordinal de origen", "ID de ubicación", "Incidencia de origen", "Ordinal de origen (GeoPol)"]
    original = row()
    original["raw"] = {name: "Original" for name in names}
    write_xlsx(
        path, [original], columns=COLUMNS, metadata=metadata(profile="source_rows", source_columns=names)
    )
    book = load_workbook(path)
    labels = [cell.value for cell in book["Datos originales"][7]]
    assert labels[1:5] == names
    assert len(labels) == len(set(labels))
    assert labels[0] == "Ordinal de origen (GeoPol 2)"
    book.close()


def test_long_metadata_keeps_overflow_sheet_last_and_all_values_recoverable(tmp_path):
    path = tmp_path / "metadata.xlsx"
    stats = write_xlsx(
        path, [row()], columns=COLUMNS, metadata=metadata(profile="source_rows", run_name="m" * 33000)
    )
    book = load_workbook(path)
    assert book.sheetnames == ["Resultados", "Datos originales", "Textos extensos"]
    assert stats["overflow"]["cells"] == 2
    assert all(book[name]["A2"].value.startswith("Texto completo") for name in book.sheetnames[:2])
    book.close()


def test_original_high_precision_numbers_survive_excel_round_trip(tmp_path):
    path = tmp_path / "precision.xlsx"
    values = [123456789012345.67, 1.2345678901234567, 1e16, -1.2345678901234567, 12.345]
    names = [f"valor{i}" for i in range(len(values))]
    original = row()
    original["raw"] = dict(zip(names, values))
    stats = write_xlsx(
        path, [original], columns=COLUMNS, metadata=metadata(profile="source_rows", source_columns=names)
    )
    book = load_workbook(path)
    for index, value in enumerate(values[:-1], 2):
        cell = book["Datos originales"].cell(8, index)
        assert cell.value == repr(value)
        assert cell.data_type == "s"
    assert book["Datos originales"]["F8"].value == 12.345
    assert book["Datos originales"]["F8"].data_type == "n"
    assert result_cells(book, stats)["latitude"].data_type == "n"
    assert result_cells(book, stats)["longitude"].data_type == "n"
    book.close()

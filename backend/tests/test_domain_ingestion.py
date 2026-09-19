import csv
import tracemalloc
import zipfile

import pytest
from openpyxl import Workbook

from geopol.domain import ingestion
from geopol.domain.ingestion import IngestionError, inspect_file, iter_records


def test_csv_preserves_malformed_recoverable_rows_and_logical_ordinals(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text(
        'ID_DENUNCIA,UBICACION\n1,"AV LAS FLORES\n123"\n2\n3,CALLE A,extra\n\n', encoding="utf-8", newline=""
    )
    records = list(iter_records(path, path.name))
    assert [r[0] for r in records] == [1, 2, 3, 4]
    assert records[0][1]["UBICACION"] == "AV LAS FLORES\n123"
    assert records[0][2] is None
    assert "COLUMNAS_DESIGUALES" in records[1][2]
    assert records[2][1]["__extra_columns__"] == ["extra"]
    assert "FILA_VACIA" in records[3][2]


@pytest.mark.parametrize("header", ["UBICACION,UBICACION", "Ubicación,ubicacion", "UBICACION,", ""])
def test_ambiguous_or_empty_headers_rejected(tmp_path, header):
    path = tmp_path / "source.csv"
    path.write_text(header + "\n", encoding="utf-8")
    with pytest.raises(IngestionError):
        list(iter_records(path, path.name))


def test_unclosed_quoted_record_raises_boundary_error(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text('UBICACION\n"no cerrar\n', encoding="utf-8")
    with pytest.raises(IngestionError, match="delimitar"):
        list(iter_records(path, path.name))


def test_csv_field_and_logical_record_limits_are_enforced(tmp_path, monkeypatch):
    path = tmp_path / "source.csv"
    path.write_text('UBICACION\n"12345678\n12345678\n12345678"\n', encoding="utf-8")
    monkeypatch.setattr(ingestion, "MAX_RECORD_CHARS", 20)
    with pytest.raises(IngestionError, match="demasiado extenso"):
        list(iter_records(path, path.name))


def test_profile_is_bounded_and_excludes_personal_fields(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text(
        "ID_DENUNCIA;UBICACION;NOMBRE;EDAD\n" + "1;CALLE A;PERSONA SINTETICA;30\n" * 30, encoding="utf-8-sig"
    )
    profile = inspect_file(path, path.name)
    assert profile["delimiter"] == ";"
    assert len(profile["sample"]) == 5
    assert all(set(row) == {"location_original"} for row in profile["sample"])
    assert "NOMBRE" in profile["columns"]


def test_xlsx_streams_selectable_sheet_and_preserves_formula_as_issue(tmp_path):
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Datos"
    worksheet.append(["ID_DENUNCIA", "UBICACION", "UBIGEO_HECHO"])
    worksheet.append(["D1", "CALLE A 2", "150101"])
    worksheet.append(["D2", '=CONCAT("CALLE", " B")', "150101"])
    workbook.create_sheet("Otra").append(["UBICACION"])
    workbook.save(path)
    rows = list(iter_records(path, path.name, sheet="Datos"))
    assert len(rows) == 2
    assert rows[0][1]["ID_DENUNCIA"] == "D1"
    assert rows[1][1]["UBICACION"].startswith("=CONCAT")
    assert "FORMULA_XLSX_NO_EVALUADA" in rows[1][2]
    assert inspect_file(path, path.name)["sheets"] == ["Datos", "Otra"]
    with pytest.raises(IngestionError, match="No existe"):
        list(iter_records(path, path.name, sheet="Ausente"))


def test_xlsx_zip_expansion_limit_checked_before_openpyxl(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", "A" * 100)
    monkeypatch.setattr(ingestion, "MAX_XLSX_SHARED_STRINGS", 20)
    with pytest.raises(IngestionError, match="cadenas compartidas"):
        list(iter_records(path, path.name))


def test_xlsx_uploaded_with_opaque_storage_name_is_read_by_declared_format(tmp_path):
    path = tmp_path / "opaque-upload-uuid"
    workbook = Workbook()
    workbook.active.append(["UBICACION"])
    workbook.active.append(["CALLE SINTETICA 10"])
    workbook.save(path)
    assert inspect_file(path, "original.xlsx")["columns"] == ["UBICACION"]
    assert list(iter_records(path, "original.xlsx"))[0][1]["UBICACION"] == "CALLE SINTETICA 10"


def test_csv_memory_does_not_grow_with_number_of_rows(tmp_path):
    path = tmp_path / "large.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["ID_DENUNCIA", "UBICACION"])
        for index in range(25_000):
            writer.writerow([index, "AV LAS FLORES 123 " + "X" * 128])
    tracemalloc.start()
    count = sum(1 for _ in iter_records(path, path.name))
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert count == 25_000
    assert peak < 2_000_000  # whole source is > 3.5 MiB; reader stays small.

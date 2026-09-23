"""Download contracts for readable Excel exports and historical GIS CSV files."""

import hashlib
import io
import json

import pytest
from openpyxl import load_workbook

from geopol.models import Export, Job, Run, Upload
from geopol.serialization import export_format
from geopol.storage import storage_file
from test_api import harness as harness


def create(harness, run_id, profile="locations", role="admin", **fields):
    response = harness.client.post(
        f"/api/runs/{run_id}/exports",
        json={"profile": profile, "format": "xlsx", **fields},
        headers=harness.headers[role],
    )
    assert response.status_code == 201, response.text
    return response.json()


def download(harness, export_id, role="admin"):
    info = harness.client.get(f"/api/exports/{export_id}", headers=harness.headers[role]).json()
    assert info["status"] == "COMPLETED", info
    response = harness.client.get(f"/api/exports/{export_id}/download", headers=harness.headers[role])
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert info["filename"].endswith(".xlsx")
    assert info["filename"] in response.headers["content-disposition"]
    assert hashlib.sha256(response.content).hexdigest() == info["sha256"]
    manifest = harness.client.get(f"/api/exports/{export_id}/manifest", headers=harness.headers[role]).json()
    return load_workbook(io.BytesIO(response.content), data_only=False), info, manifest


def records(workbook, manifest, name="Resultados"):
    sheet = workbook[name]
    details = next(item for item in manifest["workbook"]["sheets"] if item["name"] == name)
    header = details["header_row"]
    labels = [cell.value for cell in sheet[header]]
    rows = [dict(zip(labels, row)) for row in sheet.iter_rows(min_row=header + 1, values_only=True)]
    assert len(rows) == details["data_rows"]
    return rows


def label(manifest, key):
    return next(item["label"] for item in manifest["workbook"]["columns"] if item["key"] == key)


def test_xlsx_download_keeps_snapshot_coordinates_and_restricted_columns(harness):
    run = harness.run()
    result = next(item for item in harness.results(run["id"]) if item["complaint_id"] == "DEMO-A")
    accepted = harness.decide(
        result,
        "manual_point",
        latitude=-12.04,
        longitude=-77.03,
        precision="PUERTA",
        evidence="Punto sintético documentado para probar exportación Excel",
    )
    assert accepted.status_code == 200
    export = create(harness, run["id"], role="analyst")
    assert export["format"] == "xlsx"
    assert harness.decide(accepted.json(), "reopen").status_code == 200
    assert harness.worker.work_once()
    workbook, info, manifest = download(harness, export["id"], "analyst")
    assert info["row_count"] == manifest["expected_rows"] == 2
    assert manifest["format"] == "xlsx"
    assert manifest["schema_version"] == 3
    assert manifest["workbook_schema_version"] == 1
    assert manifest["sha256"] == info["sha256"]
    assert "Datos originales" not in workbook.sheetnames
    rows = records(workbook, manifest)
    resolved = next(row for row in rows if row[label(manifest, "complaint_id")] == "DEMO-A")
    unresolved = next(row for row in rows if row[label(manifest, "complaint_id")] == "DEMO-B")
    assert resolved[label(manifest, "latitude")] == -12.04
    assert resolved[label(manifest, "longitude")] == -77.03
    assert resolved[label(manifest, "revision")] == accepted.json()["revision"]
    assert unresolved[label(manifest, "latitude")] is None
    assert unresolved[label(manifest, "longitude")] is None
    text = json.dumps([list(sheet.values) for sheet in workbook], ensure_ascii=False)
    assert "Persona sintetica" not in text
    assert "person_name" not in text
    workbook.close()


def test_xlsx_source_rows_preserve_duplicates_extras_ids_and_literal_formulas(harness):
    data = (
        "complaint_id,location_original,ubigeo,=1+1\n"
        "00012345678901234567890,AV DEMOSTRACION 120,010101,=HYPERLINK(1)\n"
        "00012345678901234567890,AV DEMOSTRACION 120,010101,+SUM(1)\n"
        "DEMO-X,OTRA UBICACION,010101,@SUM(1),EXTRA\n"
    ).encode()
    run = harness.run(data)
    export = create(harness, run["id"], "source_rows", "operator", safe_spreadsheet=False)
    assert harness.worker.work_once()
    workbook, info, manifest = download(harness, export["id"], "operator")
    assert info["row_count"] == 3
    assert len(records(workbook, manifest)) == 3
    original = records(workbook, manifest, "Datos originales")
    assert [row["complaint_id"] for row in original] == [
        "00012345678901234567890",
        "00012345678901234567890",
        "DEMO-X",
    ]
    assert [row["ubigeo"] for row in original] == ["010101"] * 3
    assert [row["=1+1"] for row in original] == ["=HYPERLINK(1)", "+SUM(1)", "@SUM(1)"]
    assert json.loads(original[-1]["__extra_columns__"]) == ["EXTRA"]
    assert manifest["safe_spreadsheet"] is True
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


@pytest.mark.parametrize("role", ["reviewer", "analyst"])
def test_xlsx_original_rows_permission_applies_to_create_and_all_download_routes(harness, role):
    run = harness.run()
    rejected = harness.client.post(
        f"/api/runs/{run['id']}/exports",
        json={"format": "xlsx", "profile": "source_rows"},
        headers=harness.headers[role],
    )
    assert rejected.status_code == 403
    export = create(harness, run["id"], "source_rows")
    assert harness.worker.work_once()
    for suffix in ("", "/download", "/manifest"):
        response = harness.client.get(f"/api/exports/{export['id']}{suffix}", headers=harness.headers[role])
        assert response.status_code == 403
        assert harness.client.get(f"/api/exports/{export['id']}{suffix}").status_code == 401


def test_xlsx_metadata_is_captured_at_creation(harness):
    run = harness.run()
    export = create(harness, run["id"], "source_rows")
    with harness.sessions() as db:
        item = db.get(Run, run["id"])
        item.name = "Nombre cambiado después de solicitar"
        item.config = {**item.config, "source_columns": ["columna posterior"]}
        db.get(Upload, item.upload_id).filename = "nombre-posterior.csv"
        db.commit()
    assert harness.worker.work_once()
    workbook, _, manifest = download(harness, export["id"])
    assert manifest["run_name"] == "Prueba sintética"
    assert manifest["filename"] == "synthetic.csv"
    assert "complaint_id" in manifest["source_columns"]
    original = records(workbook, manifest, "Datos originales")
    assert original[0]["complaint_id"] == "DEMO-A"
    text = json.dumps([list(sheet.values) for sheet in workbook], ensure_ascii=False)
    assert "Nombre cambiado después" not in text
    assert "nombre-posterior.csv" not in text
    assert "columna posterior" not in text
    workbook.close()


def test_csv_is_default_and_pending_historical_manifest_remains_csv(harness):
    run = harness.run()
    export_id = harness.export(run["id"])
    with harness.sessions() as db:
        item = db.get(Export, export_id)
        legacy = {key: value for key, value in item.manifest.items() if key != "format"}
        item.manifest = {**legacy, "schema_version": 1}
        db.commit()
    assert harness.worker.work_once()
    rows = harness.download_export(export_id)
    assert len(rows) == 2
    assert "GEOPOL_review_status" not in rows[0]
    assert "GEOPOL_geometry" not in rows[0]
    info = harness.client.get(f"/api/exports/{export_id}", headers=harness.headers["admin"]).json()
    assert info["format"] == "csv"
    assert info["filename"].endswith(".csv")


def test_unknown_export_format_is_rejected(harness):
    run = harness.run()
    response = harness.client.post(
        f"/api/runs/{run['id']}/exports", json={"format": "../xlsx"}, headers=harness.headers["admin"]
    )
    assert response.status_code == 422
    with pytest.raises(ValueError):
        export_format(Export(manifest={"format": "../xlsx"}))


def test_failed_xlsx_removes_lease_temporary_file_and_does_not_publish(harness, monkeypatch):
    run = harness.run()
    export = create(harness, run["id"])

    def broken(path, rows, **kwargs):
        next(rows)
        path.write_bytes(b"unpublished workbook")
        raise ValueError("sensitive payload must not enter errors")

    monkeypatch.setattr(harness.worker, "write_xlsx", broken)
    assert harness.worker.work_once()
    info = harness.client.get(f"/api/exports/{export['id']}", headers=harness.headers["admin"]).json()
    assert info["status"] == "FAILED"
    assert "sensitive payload" not in info["error"]
    destination = storage_file("exports", export["id"], ".xlsx")
    assert not destination.exists()
    assert not list(destination.parent.glob(f"{export['id']}.*.part"))


def test_expired_xlsx_lease_cannot_publish(harness, monkeypatch):
    run = harness.run()
    export = create(harness, run["id"])
    original_writer = harness.worker.write_xlsx

    def expire_after_writing(path, rows, **kwargs):
        stats = original_writer(path, rows, **kwargs)
        with harness.sessions() as db:
            job = db.query(Job).filter_by(target_id=export["id"]).one()
            job.lease_until = 0
            db.commit()
        return stats

    monkeypatch.setattr(harness.worker, "write_xlsx", expire_after_writing)
    assert harness.worker.work_once()
    destination = storage_file("exports", export["id"], ".xlsx")
    assert not destination.exists()
    assert not list(destination.parent.glob(f"{export['id']}.*.part"))
    response = harness.client.get(f"/api/exports/{export['id']}/download", headers=harness.headers["admin"])
    assert response.status_code == 409


def test_excel_limit_gives_safe_actionable_error(harness, monkeypatch):
    from geopol import excel_export

    run = harness.run()
    export = create(harness, run["id"])
    monkeypatch.setattr(excel_export, "MAX_ROWS", excel_export.HEADER_ROW)
    assert harness.worker.work_once()
    info = harness.client.get(f"/api/exports/{export['id']}", headers=harness.headers["admin"]).json()
    assert info["status"] == "FAILED"
    assert "Utilice CSV" in info["error"]
    assert "DEMO" not in info["error"]
    assert not storage_file("exports", export["id"], ".xlsx").exists()

"""Streaming presentation of an export snapshot as an Excel workbook.

The writer neither changes geographic decisions nor accesses persistence. All
strings are literal spreadsheet values, including original source data. Excel's
cell limit is handled with recoverable fragments, never by truncating content.
"""

import json
import math
import re
from collections.abc import Iterable, Mapping, Sequence
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_ROWS = 1_048_576
MAX_COLUMNS = 16_384
MAX_CELL_UNITS = 32_767
FRAGMENT_UNITS = 30_000
HEADER_ROW = 7
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

_FIELDS = {
    "complaint_id": ("Denuncia", 20),
    "ubigeo": ("UBIGEO", 12),
    "quality_flag": ("Flag de calidad", 18),
    "source_quality_flag": ("Flag de origen", 18),
    "source_quality_flag_original": ("Valor original del flag", 23),
    "review_state": ("Estado de revisión", 24),
    "quality_flag_reason": ("Motivo del flag", 56),
    "quality_code": ("Calidad (provisional)", 22),
    "quality_stage": ("Etapa de evaluación", 23),
    "quality_status": ("Estado de calidad", 23),
    "quality_reason": ("Motivo de calidad", 56),
    "quality_policy_version": ("Versión de calidad", 22),
    "resolution": ("Resolución", 29),
    "location_normalized": ("Dirección normalizada", 56),
    "product": ("Producto geográfico", 25),
    "precision": ("Precisión espacial", 22),
    "latitude": ("Latitud", 16),
    "longitude": ("Longitud", 16),
    "review_bucket": ("Tipo de pendiente", 23),
    "reason": ("Motivo y evidencia", 64),
    "review_status": ("Estado de revisión", 21),
    "method": ("Método", 25),
    "evidence_band": ("Nivel de evidencia", 22),
    "source_row_count": ("Filas de origen", 18),
    "location_original": ("Dirección original", 56),
    "revision": ("Revisión", 12),
    "id": ("ID de ubicación", 39),
    "geometry": ("Geometría GeoJSON", 64),
}
_TEXT_FIELDS = {"id", "complaint_id", "ubigeo"}
_TECHNICAL_FIELDS = {"id", "geometry"}

_FONT = Font(name="Arial", size=10, color="183A52")
_TITLE = Font(name="Arial", size=14, bold=True, color="12426B")
_META = Font(name="Arial", size=10, color="48657B")
_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="12426B")
_BAND_FILL = PatternFill("solid", fgColor="F0F6FB")
_WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
_TITLE_ALIGN = Alignment(vertical="center")
_ACCENT = Border(bottom=Side(style="thin", color="008FD3"))
_STATE_COLORS = {
    "ACEPTADO_AUTOMATICO": ("E5F2EC", "176547"),
    "ACEPTADO_MANUAL": ("E9F3FC", "12426B"),
    "EXCLUIDO_FLAG_10": ("F2F4F7", "526779"),
    "REVISION_REQUERIDA": ("FFF3DC", "805700"),
    "SIN_COINCIDENCIA": ("F4F5F7", "566473"),
    "NO_EVALUABLE_REFERENCIA": ("FFF3DC", "805700"),
    "ERROR": ("FDECEC", "9E3030"),
}
_STATE_STYLES = {
    key: (PatternFill("solid", fgColor=fill), Font(name="Arial", size=10, bold=True, color=color))
    for key, (fill, color) in _STATE_COLORS.items()
}


class ExcelExportError(ValueError):
    """Actionable workbook errors whose messages contain no source data."""


def _text_units(value: str) -> int:
    return len(value.encode("utf-16-le")) // 2


def _fragments(value: str):
    """Use UTF-16 units because a supplementary character occupies two in Excel."""
    start = 0
    units = 0
    for index, char in enumerate(value):
        width = 2 if ord(char) > 0xFFFF else 1
        if units + width > FRAGMENT_UNITS:
            yield value[start:index]
            start, units = index, 0
        units += width
    if start < len(value):
        yield value[start:]


def _literal(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise ExcelExportError(
            "Excel no admite valores numéricos no finitos. Revise los datos o utilice CSV."
        )
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


def _unique_label(label, used):
    candidate = label
    suffix = 1
    while candidate in used:
        candidate = f"{label} (GeoPol{f' {suffix}' if suffix > 1 else ''})"
        suffix += 1
    used.add(candidate)
    return candidate


def _source_value(value):
    # Excel has 15 significant decimal digits. Preserve source measurements as
    # text when numeric serialization could round them; geographic output
    # coordinates retain their numeric type and dedicated display format.
    if isinstance(value, float) and math.isfinite(value):
        representation = repr(value)
        mantissa = representation.lower().split("e", 1)[0]
        significant = mantissa.lstrip("+-").replace(".", "").lstrip("0")
        if len(significant) > 15 or value.is_integer() and abs(value) >= 10**15:
            return representation
    return value


class _ExcelWriter:
    def __init__(self):
        self.book = Workbook(write_only=True)
        self.book.properties.creator = "GeoPol"
        self.book.properties.title = "Resultados de normalización y geocodificación"
        self.sheets = []
        self.overflow = None
        self.overflow_cells = 0
        self.overflow_segments = 0

    def _cell(self, sheet, value, row, key, *, font=_FONT, fill=None, centered=False):
        value = _literal(value)
        if isinstance(value, str):
            if _INVALID_XML.search(value):
                raise ExcelExportError(
                    f"Excel no admite un carácter de control en {sheet.title}, fila {row}. "
                    "Utilice CSV para conservar el valor original."
                )
            if _text_units(value) > MAX_CELL_UNITS:
                value = self._long_text(sheet.title, row, key, value)
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            # Explicit type avoids formulas and preserves leading zeroes verbatim.
            cell.data_type = "s"
            cell.number_format = "@"
        elif isinstance(value, int):
            cell.number_format = "#,##0"
        cell.font = font
        cell.alignment = _CENTER if centered else _WRAP
        if fill is not None:
            cell.fill = fill
        return cell

    def _long_text(self, sheet_name, row, key, value):
        if self.overflow is None:
            self.overflow = self.book.create_sheet("Textos extensos")
            self._configure(
                self.overflow,
                [
                    ("Clave", 18),
                    ("Hoja", 24),
                    ("Fila Excel", 14),
                    ("Campo", 30),
                    ("Parte", 12),
                    ("Total partes", 14),
                    ("Texto original", 100),
                ],
                header_row=3,
            )
            self.overflow.merged_cells.add("A1:G1")
            self.overflow.merged_cells.add("A2:G2")
            self.overflow.row_dimensions[2].height = 20
            self.overflow.append(
                [
                    self._cell(
                        self.overflow, "Textos que exceden el límite de una celda", 1, "título", font=_TITLE
                    )
                ]
            )
            self.overflow.append(
                [
                    self._cell(
                        self.overflow,
                        "Para recuperar el valor, una las partes de cada clave en orden, sin separadores.",
                        2,
                        "nota",
                        font=_META,
                    )
                ]
            )
            self._headers(
                self.overflow,
                ["Clave", "Hoja", "Fila Excel", "Campo", "Parte", "Total partes", "Texto original"],
                3,
            )
        self.overflow_cells += 1
        token = f"T{self.overflow_cells:08d}"
        # Count without retaining another full copy of potentially large geometries.
        count = sum(1 for _ in _fragments(value))
        if self.overflow_segments + count + 3 > MAX_ROWS:
            raise ExcelExportError("Los textos extensos exceden el límite de filas de Excel. Utilice CSV.")
        for part, fragment in enumerate(_fragments(value), 1):
            target_row = self.overflow_segments + 4
            values = [token, sheet_name, row, key, part, count, fragment]
            cells = [
                self._cell(self.overflow, item, target_row, name)
                for name, item in zip(("clave", "hoja", "fila", "campo", "parte", "partes", "texto"), values)
            ]
            self.overflow.append(cells)
            self.overflow_segments += 1
        return f"Texto completo en Textos extensos: {token} ({count} partes)"

    def _configure(self, sheet, fields, *, header_row=HEADER_ROW, hidden_start=None):
        if len(fields) > MAX_COLUMNS:
            raise ExcelExportError("La exportación excede el límite de columnas de Excel. Utilice CSV.")
        sheet.sheet_view.showGridLines = False
        sheet.sheet_format.defaultRowHeight = 42
        sheet.sheet_view.zoomScale = 90
        sheet.freeze_panes = f"C{header_row + 1}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.outlinePr.summaryRight = True
        sheet.sheet_properties.tabColor = "12426B" if sheet.title == "Resultados" else "008FD3"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f"1:{header_row}"
        for index, (_, width) in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        if hidden_start is not None:
            sheet.column_dimensions.group(
                get_column_letter(hidden_start), get_column_letter(len(fields)), outline_level=1, hidden=True
            )
            if len(fields) < MAX_COLUMNS:
                sheet.column_dimensions[get_column_letter(len(fields) + 1)].collapsed = True
        sheet.oddFooter.center.text = "GeoPol · Página &P de &N"
        sheet.row_dimensions[1].height = 25

    def _headers(self, sheet, labels, row=HEADER_ROW):
        sheet.row_dimensions[row].height = 34
        cells = []
        for label in labels:
            cell = self._cell(sheet, label, row, "encabezado", font=_HEADER, fill=_HEADER_FILL)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = _ACCENT
            cells.append(cell)
        sheet.append(cells)

    def make_sheet(self, name, fields, metadata, *, hidden_start=None):
        sheet = self.book.create_sheet(name)
        self._configure(sheet, fields, hidden_start=hidden_start)
        reference = metadata.get("reference")
        reference_text = "Sin catálogo de referencia"
        if reference:
            reference_text = str(reference.get("name") or reference.get("id") or "Catálogo de referencia")
            if reference.get("version"):
                reference_text += f". Versión: {reference['version']}"
        source_profile = metadata.get("profile") == "source_rows"
        notes = (
            "Una fila por registro de origen. El ordinal vincula ambas hojas; una ubicación puede repetirse."
            if source_profile
            else "Una fila por ubicación. Áreas y tramos conservan su geometría; no representan una puerta."
        )
        if name == "Datos originales":
            notes = "Valores de origen conservados. Ordinal e ID de ubicación permiten vincularlos con Resultados."
        lines = [
            "GeoPol · Resultados" if name == "Resultados" else "GeoPol · Datos originales",
            f"Procesamiento: {metadata.get('run_name') or 'Sin nombre'}. Archivo: {metadata.get('filename') or '—'}",
            f"Instantánea: {metadata.get('snapshot_at') or '—'}. Reglas: {metadata.get('rules_version') or '—'}. "
            f"Filas: {metadata.get('expected_rows', '—')}",
            f"Referencia: {reference_text}",
            notes,
            None,
        ]
        merge_end = get_column_letter(max(1, min(len(fields), 10)))
        for index, text in enumerate(lines, 1):
            if text is not None:
                if merge_end != "A":
                    sheet.merged_cells.add(f"A{index}:{merge_end}{index}")
                cell = self._cell(sheet, text, index, "metadatos", font=_TITLE if index == 1 else _META)
                cell.alignment = _TITLE_ALIGN if index == 1 else _WRAP
                sheet.row_dimensions[index].height = 25 if index == 1 else 20
                sheet.append([cell])
            else:
                sheet.row_dimensions[index].height = 8
                sheet.append([])
        self._headers(sheet, [label for label, _ in fields])
        self.sheets.append(
            {"name": name, "data_rows": 0, "header_row": HEADER_ROW, "column_count": len(fields)}
        )
        return sheet

    def data_row(self, sheet, values, keys, index, *, resolution=None):
        excel_row = HEADER_ROW + index
        if excel_row > MAX_ROWS:
            raise ExcelExportError("La exportación excede el límite de filas de Excel. Utilice CSV.")
        fill = _BAND_FILL if index % 2 == 0 else _WHITE_FILL
        cells = []
        for key, value in zip(keys, values):
            if key in _TEXT_FIELDS and value is not None:
                value = str(value)
            cell = self._cell(
                sheet,
                value,
                excel_row,
                key,
                fill=fill,
                centered=key in {"ubigeo", "revision", "source_row_count", "ordinal"},
            )
            if key in {"latitude", "longitude"} and isinstance(value, (int, float)):
                cell.number_format = "0.00000000"
            if key == "resolution" and resolution in _STATE_STYLES:
                cell.fill, cell.font = _STATE_STYLES[resolution]
            cells.append(cell)
        sheet.append(cells)

    def save(self, path, count):
        for stats in self.sheets:
            stats["data_rows"] = count
            sheet = self.book[stats["name"]]
            sheet.auto_filter.ref = (
                f"A{HEADER_ROW}:{get_column_letter(stats['column_count'])}{HEADER_ROW + count}"
            )
        if self.overflow is not None:
            self.book.move_sheet(
                self.overflow.title, offset=len(self.book.sheetnames) - self.book.index(self.overflow) - 1
            )
            self.overflow.auto_filter.ref = f"A3:G{3 + self.overflow_segments}"
            self.sheets.append(
                {
                    "name": "Textos extensos",
                    "data_rows": self.overflow_segments,
                    "header_row": 3,
                    "column_count": 7,
                }
            )
        self.book.save(path)

    def cleanup(self):
        # Write-only worksheets use temporary XML files. On an aborted export,
        # close and remove them as save() normally does; workers remain long-lived.
        for sheet in self.book:
            if not sheet.closed:
                sheet.close()
            writer = sheet._writer
            if writer is not None and isinstance(writer.out, str) and Path(writer.out).exists():
                writer.cleanup()
        self.book.close()


def write_xlsx(
    path: str | Path, rows: Iterable[Mapping], *, columns: Sequence[str], metadata: Mapping
) -> dict:
    """Write one pass of immutable rows; return aggregate, JSON-safe file metadata."""
    if not columns or len(set(columns)) != len(columns):
        raise ExcelExportError("La exportación requiere columnas únicas y no vacías.")
    profile = metadata.get("profile", "locations")
    if profile not in {"locations", "source_rows"}:
        raise ExcelExportError("Perfil de exportación de Excel no reconocido.")
    expected = metadata.get("expected_rows")
    if expected is not None and (not isinstance(expected, int) or expected < 0):
        raise ExcelExportError("La cantidad esperada de filas debe ser un entero no negativo.")
    if expected is not None and expected + HEADER_ROW > MAX_ROWS:
        raise ExcelExportError("La exportación excede el límite de filas de Excel. Utilice CSV.")
    # Known presentation order first; any future schema columns remain preserved.
    ordered = [key for key in _FIELDS if key in columns and key not in _TECHNICAL_FIELDS]
    ordered += [key for key in columns if key not in _FIELDS]
    ordered += [key for key in _FIELDS if key in columns and key in _TECHNICAL_FIELDS]
    result_keys = (["ordinal"] if profile == "source_rows" else []) + ordered
    fields = [("Ordinal de origen", 18)] if profile == "source_rows" else []
    fields += [
        ("Bandeja de revisión", 23)
        if key == "review_status" and "review_state" in columns
        else _FIELDS.get(key, (key, 25))
        for key in ordered
    ]
    source_columns = list(metadata.get("source_columns") or [])
    if len(fields) > MAX_COLUMNS or profile == "source_rows" and len(source_columns) + 3 > MAX_COLUMNS:
        raise ExcelExportError("La exportación excede el límite de columnas de Excel. Utilice CSV.")
    technical_count = sum(key in _TECHNICAL_FIELDS for key in ordered)
    writer = _ExcelWriter()
    count = 0
    try:
        results = writer.make_sheet(
            "Resultados",
            fields,
            metadata,
            hidden_start=len(fields) - technical_count + 1 if technical_count else None,
        )
        originals = None
        if profile == "source_rows":
            used_labels = {str(key) for key in source_columns}
            original_fields = [(_unique_label("Ordinal de origen", used_labels), 18)]
            original_fields += [(str(key), 25) for key in source_columns]
            original_fields += [
                (_unique_label("ID de ubicación", used_labels), 39),
                (_unique_label("Incidencia de origen", used_labels), 50),
            ]
            originals = writer.make_sheet("Datos originales", original_fields, metadata)
        for record in rows:
            count += 1
            snapshot = record.get("snapshot") or {}
            result_values = [record.get("ordinal")] if profile == "source_rows" else []
            result_values += [snapshot.get(key) for key in ordered]
            writer.data_row(results, result_values, result_keys, count, resolution=snapshot.get("resolution"))
            if originals is not None:
                raw = record.get("raw") or {}
                values = [record.get("ordinal")] + [_source_value(raw.get(key)) for key in source_columns]
                values += [snapshot.get("id"), record.get("issue")]
                # Original field names can equal semantic keys. Prefixing prevents
                # the result-column presentation rules from changing source types.
                keys = ["ordinal"] + [f"origen:{key}" for key in source_columns] + ["id", "issue"]
                writer.data_row(originals, values, keys, count)
        if expected is not None and count != expected:
            raise ExcelExportError("La exportación no coincide con la cardinalidad de la instantánea.")
        writer.save(path, count)
        return {
            "row_count": count,
            "sheets": writer.sheets,
            "columns": [{"key": key, "label": label} for key, (label, _) in zip(result_keys, fields)],
            "overflow": {"cells": writer.overflow_cells, "segments": writer.overflow_segments},
        }
    finally:
        writer.cleanup()

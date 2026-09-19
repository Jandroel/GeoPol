"""Streaming local CSV/XLSX readers with explicit limits and recoverable row errors."""

from __future__ import annotations

import csv
import datetime as dt
import io
import itertools
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from .normalization import column_key, suggest_mapping, text

MAX_COLUMNS = 512
MAX_FIELD_CHARS = 1_048_576
MAX_RECORD_CHARS = 8_388_608
MAX_XLSX_UNCOMPRESSED = 536_870_912
MAX_XLSX_ENTRIES = 4096
MAX_XLSX_SHARED_STRINGS = 67_108_864
_SUPPORTED_ENCODINGS = {"utf-8", "utf-8-sig", "cp1252", "latin-1"}


class IngestionError(ValueError):
    """The source cannot safely be parsed without losing record boundaries."""


def _extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise IngestionError("Formato no admitido. Use CSV o XLSX; XLS y archivos con macros no están admitidos.")
    return extension


def _headers(row: tuple | list) -> list[str]:
    if not row or all(v is None or text(v) == "" for v in row):
        raise IngestionError("La hoja o archivo no contiene encabezados.")
    if len(row) > MAX_COLUMNS:
        raise IngestionError(f"Límite de columnas excedido: máximo {MAX_COLUMNS}.")
    columns = [text(v) for v in row]
    if any(not c for c in columns):
        raise IngestionError("Encabezado vacío: asigne un nombre a cada columna antes de importar.")
    keys = [column_key(c) for c in columns]
    if len(set(keys)) != len(keys):
        raise IngestionError("Encabezados duplicados o equivalentes; el mapeo sería ambiguo.")
    return columns


class _BoundedLines:
    """Bound both physical lines and quoted records while csv handles logical rows."""

    def __init__(self, stream: io.TextIOBase):
        self.stream = stream
        self.record_chars = 0

    def __iter__(self):
        return self

    def __next__(self):
        remaining = MAX_RECORD_CHARS - self.record_chars
        line = self.stream.readline(remaining + 1)
        if not line:
            raise StopIteration
        self.record_chars += len(line)
        if self.record_chars > MAX_RECORD_CHARS:
            raise IngestionError(f"Registro CSV demasiado extenso; límite {MAX_RECORD_CHARS} caracteres.")
        return line


@contextmanager
def _csv_reader(path: Path, delimiter: str, encoding: str):
    if encoding not in _SUPPORTED_ENCODINGS:
        raise IngestionError("Codificación no admitida: seleccione utf-8-sig, utf-8, cp1252 o latin-1.")
    if len(delimiter) != 1 or delimiter in {'"', "\r", "\n", "\0"}:
        raise IngestionError("El separador debe ser un solo carácter distinto de comillas o salto de línea.")
    csv.field_size_limit(MAX_FIELD_CHARS)
    try:
        with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
            lines = _BoundedLines(stream)
            reader = csv.reader(lines, delimiter=delimiter, strict=True)
            yield reader, lines
    except (UnicodeError, csv.Error) as exc:
        raise IngestionError(f"No se puede delimitar con seguridad el CSV: {exc}") from exc


def _xlsx_safety(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES or sum(info.file_size for info in entries) > MAX_XLSX_UNCOMPRESSED:
                raise IngestionError("XLSX excede el límite de 512 MiB descomprimidos o 4096 componentes.")
            for info in entries:
                if info.flag_bits & 1:
                    raise IngestionError("No se admiten libros cifrados.")
                if info.filename == "xl/sharedStrings.xml" and info.file_size > MAX_XLSX_SHARED_STRINGS:
                    raise IngestionError("XLSX supera 64 MiB de cadenas compartidas; conviértalo a CSV para procesarlo en flujo.")
    except zipfile.BadZipFile as exc:
        raise IngestionError("El archivo no es un XLSX válido.") from exc


@contextmanager
def _workbook(path: Path):
    _xlsx_safety(path)
    # Uploaded blobs deliberately use opaque names without an extension.
    with path.open("rb") as stream:
        try:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        except Exception as exc:
            raise IngestionError(f"No se pudo abrir el XLSX: {type(exc).__name__}.") from exc
        try:
            yield workbook
        finally:
            workbook.close()


def _worksheet(workbook, sheet: str | None):
    if sheet is not None and sheet not in workbook.sheetnames:
        raise IngestionError(f"No existe la hoja solicitada: {sheet}.")
    if not workbook.sheetnames:
        raise IngestionError("El libro no contiene hojas.")
    selected = workbook[sheet or workbook.sheetnames[0]]
    # Ignore misleading exported dimension metadata; actual cells determine limits.
    selected.reset_dimensions()
    return selected


def _serializable(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def _row(columns: list[str], values: list, formulas: bool = False) -> tuple[dict, str | None]:
    issues = []
    if len(values) > MAX_COLUMNS:
        raise IngestionError(f"Fila excede el límite de {MAX_COLUMNS} columnas; procesamiento detenido sin omitir filas.")
    if len(values) != len(columns):
        issues.append(f"COLUMNAS_DESIGUALES: esperadas {len(columns)}, recibidas {len(values)}")
    result = {name: _serializable(values[i]) if i < len(values) else None for i, name in enumerate(columns)}
    if len(values) > len(columns):
        extra_key = "__extra_columns__"
        while extra_key in result:
            extra_key += "_"
        result[extra_key] = [_serializable(v) for v in values[len(columns):]]
    if not any(v is not None and text(v) for v in values):
        issues.append("FILA_VACIA")
    if formulas:
        issues.append("FORMULA_XLSX_NO_EVALUADA: se conserva la expresión original")
    if any(isinstance(v, str) and len(v) > MAX_FIELD_CHARS for v in values):
        issues.append("CAMPO_DEMASIADO_EXTENSO")
    return result, "; ".join(issues) or None


def iter_records(path: Path, filename: str, sheet: str | None = None, delimiter: str = ",", encoding: str = "utf-8-sig") -> Iterator[tuple[int, dict, str | None]]:
    """Yield every logical source record; ordinal 1 follows the header."""
    path = Path(path)
    if _extension(filename) == ".csv":
        with _csv_reader(path, delimiter, encoding) as (reader, lines):
            columns = _headers(next(reader, []))
            lines.record_chars = 0
            for ordinal, values in enumerate(reader, start=1):
                lines.record_chars = 0
                raw, issue = _row(columns, values)
                yield ordinal, raw, issue
    else:
        with _workbook(path) as workbook:
            selected = _worksheet(workbook, sheet)
            rows = selected.iter_rows()
            header = next(rows, ())
            columns = _headers([c.value for c in header])
            for ordinal, cells in enumerate(rows, start=1):
                values = [c.value for c in cells]
                # Streaming worksheets can omit trailing empty cells per row.
                if len(values) < len(columns):
                    values.extend([None] * (len(columns) - len(values)))
                raw, issue = _row(columns, values, any(c.data_type == "f" for c in cells))
                yield ordinal, raw, issue


def inspect_file(path: Path, filename: str, sheet: str | None = None) -> dict:
    """Read headers and at most five records; never expose personal source columns."""
    path = Path(path)
    extension = _extension(filename)
    sheets, warnings, delimiter, encoding = [], [], ",", "utf-8-sig"
    if extension == ".csv":
        raw_sample = path.open("rb")
        try:
            sample_bytes = raw_sample.read(16_384)
        finally:
            raw_sample.close()
        try:
            sample_text = sample_bytes.decode(encoding)
        except UnicodeDecodeError as exc:
            # Only retry if error is not a truncated UTF-8 sequence at sample edge.
            if exc.end == len(sample_bytes) and exc.reason == "unexpected end of data":
                sample_text = sample_bytes[:exc.start].decode(encoding)
            else:
                encoding = "cp1252"
                try:
                    sample_text = sample_bytes.decode(encoding)
                except UnicodeDecodeError as invalid:
                    raise IngestionError("Codificación desconocida; convierta el CSV a UTF-8.") from invalid
                warnings.append("CODIFICACION_SUGERIDA_CP1252: confirme antes de procesar.")
        try:
            delimiter = csv.Sniffer().sniff(sample_text, delimiters=",;\t|").delimiter
        except csv.Error:
            warnings.append("SEPARADOR_NO_DETECTADO: se propone coma; confirme el contrato.")
        with _csv_reader(path, delimiter, encoding) as (reader, _):
            columns = _headers(next(reader, []))
    else:
        with _workbook(path) as workbook:
            sheets = workbook.sheetnames
            selected = _worksheet(workbook, sheet)
            columns = _headers([c.value for c in next(selected.iter_rows(), ())])
    mapping = suggest_mapping(columns)
    records = iter_records(path, filename, sheet=sheet, delimiter=delimiter, encoding=encoding)
    samples = []
    try:
        for _, raw, issue in itertools.islice(records, 5):
            safe = {canonical: raw.get(source) for canonical, source in mapping.items() if canonical != "complaint_id"}
            samples.append(safe)
            if issue:
                warnings.append(issue)
    finally:
        records.close()
    if "location_original" not in mapping and "latitude" not in mapping:
        warnings.append("SIN_LOCALIZADOR_RECONOCIDO: configure el mapeo antes de procesar.")
    return {"columns": columns, "sheets": sheets, "suggested_mapping": mapping, "sample": samples,
            "warnings": warnings, "delimiter": delimiter, "encoding": encoding, "sheet": sheet or (sheets[0] if sheets else None)}
